"""Gera o índice de atividades para os colaboradores externos que escrevem os resumos.

Saídas (em resumos-colaboradores/): INDICE-RESUMOS.xlsx, INDICE-RESUMOS.csv e indice.json (este último é lido por
`import_agent_summaries.py`; o colaborador não precisa dele).
Cada linha = uma atividade de Teoria (= um trecho). O ID (S0001…) segue a ORDEM DE PRIORIDADE (summary_queue.py): primeiro o edital do
CFO CBMPE (Oficial), dentro dele os assuntos comuns aos 4 concursos e depois os de maior importância; por último os que só constam
em outros concursos. Os IDs são gravados na 1ª geração e preservados nas seguintes (o que for novo entra no fim).

Uso: python pipeline/build_summary_index.py
"""
from __future__ import annotations

import csv
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))
from config import CONTENT_DIR, ROOT  # noqa: E402
from summary_queue import done_ids  # noqa: E402
from textutil import aula_key_map, load_pages  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

OUT = ROOT / "resumos-colaboradores"
CONTESTS = {"CBMPE_OFICIAL": "CFO CBMPE", "CBMPE_SOLDADO": "Soldado CBMPE", "PMPE_OFICIAL": "CFO PMPE", "PMPE_SOLDADO": "Soldado PMPE"}
HEADERS = [
    "ID", "Matéria", "Aula", "Trecho", "Atividade (assuntos do trecho)", "Arquivo PDF do material", "Pág. inicial", "Pág. final",
    "Págs.", "Alvo máx. (palavras)", "No edital do CFO CBMPE", "Nº de concursos (de 4)", "Prioridade", "Lote", "Sublote (20)", "Status",
]
# Divisão do trabalho em 6 lotes por disciplina (Fase 1 = assuntos do edital do CFO CBMPE, balanceados por palavras-alvo).
# Fase 2 (fora do edital do CFO CBMPE: Atualidades, Raciocínio Lógico, História de PE) vira lote F2-A/F2-B, distribuído depois.
LOT_OF = {
    "direito-constitucional": "1", "direito-administrativo": "1",
    "informatica": "2", "lingua-inglesa": "2",
    "biologia": "3", "lingua-portuguesa": "3",
    "estatistica": "4", "direito-penal-militar": "4",
    "matematica": "5", "fisica": "5",
    "quimica": "6", "lingua-espanhola": "6",
    "atualidades": "F2-A", "raciocinio-logico": "F2-B", "historia-pe": "F2-B", "legislacoes-pe": "F2-B",
}


def main() -> None:
    queue = json.loads((CONTENT_DIR / "summaries" / "_queue.json").read_text(encoding="utf-8"))["queue"]
    importance = json.loads((CONTENT_DIR / "importance" / "segments.json").read_text(encoding="utf-8"))["segments"]
    catalog = json.loads((CONTENT_DIR / "catalog.json").read_text(encoding="utf-8"))
    seg_info: dict[str, dict] = {}
    for f in (CONTENT_DIR / "segments").glob("*.json"):
        for s in json.loads(f.read_text(encoding="utf-8"))["segments"]:
            seg_info[s["id"]] = s
    subject_name = {s["id"]: s["name"] for s in catalog["subjects"]}
    aula_info = {a["id"]: a for s in catalog["subjects"] for a in s["aulas"]}
    pdf_of = {}
    key_map = aula_key_map()
    for aula_id, key in key_map.items():
        pdf_of[aula_id] = load_pages(key)["file"]
    done = done_ids()

    # IDs estáveis: reaproveita os já atribuídos (indice.json anterior) e acrescenta os novos ao fim
    prev_file = OUT / "indice.json"
    prev = {r["trecho"]: r["id"] for r in json.loads(prev_file.read_text(encoding="utf-8"))["rows"]} if prev_file.exists() else {}
    next_n = max((int(i[1:]) for i in prev.values()), default=0) + 1
    rows = []
    for r in queue:
        sid = r["id"]
        if sid not in prev:
            prev[sid] = f"S{next_n:04d}"
            next_n += 1
        s = seg_info[sid]
        a = aula_info[s["aula"]]
        authored = a.get("source") == "authored"
        topics = s.get("topicsCovered") or [s.get("startTopic") or ""]
        number = int(a["number"])
        aula_label = f"Complemento {number - 100:02d} — {a['title']}" if authored else f"Aula {number:02d} — {a['title']}"
        pdf = (os.path.basename(a.get("materialPath", "")) if authored else pdf_of.get(s["aula"], "")) or ""
        rec = importance[sid]["contests"].get("CBMPE_OFICIAL", {})
        rows.append(
            {
                "id": prev[sid], "subject": sid.split("/")[0], "subjectName": subject_name[sid.split("/")[0]], "aula": s["aula"], "aulaLabel": aula_label,
                "trecho": sid, "activity": "; ".join(t for t in topics if t), "pdf": pdf, "startPage": s["startPage"], "endPage": s["endPage"],
                "pages": s["pages"], "targetWords": r["targetWords"], "inCbmpe": bool(rec.get("appears")), "shared": r["shared"],
                "priority": r["rank"], "authored": authored, "done": sid in done,
            }
        )
    rows.sort(key=lambda x: x["id"])
    OUT.mkdir(parents=True, exist_ok=True)

    # sublotes de 20 atividades "A fazer", na ordem de prioridade (ID), dentro de cada lote
    counter: dict[str, int] = {}
    for x in rows:
        x["lot"] = LOT_OF.get(x["subject"], "F2-B")
        if x["done"]:
            x["sublot"] = ""
            continue
        n = counter.get(x["lot"], 0)
        counter[x["lot"]] = n + 1
        x["sublot"] = f"L{x['lot']}-{n // 20 + 1:02d}" if x["lot"].isdigit() else f"{x['lot']}-{n // 20 + 1:02d}"

    (OUT / "indice.json").write_text(json.dumps({"version": 1, "rows": rows}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    def line(i: int, x: dict) -> list:
        return [
            x["id"], x["subjectName"], x["aulaLabel"], x["trecho"], x["activity"], x["pdf"], x["startPage"], x["endPage"], x["pages"], x["targetWords"],
            "Sim" if x["inCbmpe"] else "Não", x["shared"], x["priority"], x["lot"], x["sublot"],
            "Feito (já entregue)" if x["done"] else "A fazer",
        ]

    def clean(v):
        """Remove caracteres de controle (o Excel os rejeita) que vêm de títulos extraídos do PDF."""
        return re.sub("[" + "".join(chr(c) for c in list(range(0, 9)) + [11, 12] + list(range(14, 32)) + [0xFFFD]) + "]", "", v).strip() if isinstance(v, str) else v

    table = [[clean(c) for c in line(i, x)] for i, x in enumerate(rows)]
    with open(OUT / "INDICE-RESUMOS.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(HEADERS)
        w.writerows(table)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Índice"
    ws.append(HEADERS)
    for t in table:
        ws.append(t)
    head_fill = PatternFill("solid", fgColor="1F2937")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [8, 26, 38, 30, 60, 46, 9, 9, 7, 11, 11, 11, 10, 8, 11, 18]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[15].value != "A fazer":
            for c in row:
                c.fill = PatternFill("solid", fgColor="E5E7EB")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    legend = wb.create_sheet("Como usar")
    for line_ in [
        "Cada linha é UMA atividade de Teoria (um trecho de ~1 h do material). O colaborador recebe: 'execute o Sublote L3-02' (filtre a coluna Sublote) e faz as linhas na ordem do ID.",
        "IDs seguem a ordem de prioridade: edital do CFO CBMPE primeiro (comuns aos 4 concursos antes), depois os demais.",
        "Material: abra o PDF indicado na coluna F e leia SOMENTE das págs. inicial a final (colunas G–H). Para 'Complemento', o PDF é um material curto do MentorIA (leia inteiro).",
        "Alvo máx. (palavras): teto do resumo analítico (≈ 20 min de leitura). O tamanho é o que o conteúdo pedir; não encha.",
        "Linhas cinzas (Feito) já têm resumo: pule.",
        "Entrega: ver INSTRUCOES.md (um arquivo .md por atividade, nomeado ID__AGENTE.md, dentro da pasta resumos_<AGENTE>/).",
    ]:
        legend.append([line_])
    legend.column_dimensions["A"].width = 160
    wb.save(OUT / "INDICE-RESUMOS.xlsx")

    pend = [x for x in rows if not x["done"]]
    print(f"{len(rows)} atividades no índice ({len(pend)} a fazer, {len(rows) - len(pend)} feitas) -> {OUT}")
    print(f"CFO CBMPE: {sum(1 for x in pend if x['inCbmpe'])} a fazer; palavras-alvo (teto): {sum(x['targetWords'] for x in pend if x['inCbmpe']):,}")


if __name__ == "__main__":
    main()
