"""Planilha dos cadernos a montar no TecConcursos para o concurso da PCPE (Agente de Polícia).

    pipeline/.venv/Scripts/python pipeline/questoes/planilha_cadernos_pcpe.py

Gera: docs/questoes/downloads/cadernos-a-montar-pcpe.xlsx
Contém:
1. Aba 'Cadernos': Visão geral das 12 disciplinas da PCPE Agente (inclusive as 5 novas),
   pesos estimados na prova, meta de questões Cebraspe, filtros recomendados e observações.
2. Aba 'Tópicos do Edital': Lista detalhada de todos os tópicos do edital da PCPE Agente
   para cada disciplina, com indicação de aula/cobertura, questões Cebraspe já aprovadas no banco,
   meta mínima (10 por assunto/aula) e termos de busca/assuntos do TEC a marcar.
3. Aba 'Filtros e Instruções': Parâmetros exatos de busca, filtro da banca Cebraspe, faixas de anos,
   níveis de escolaridade, opções de impressão e rotina de importação.
"""
from __future__ import annotations

import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
WEB = ROOT / "web"
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# Configuração dos 12 Cadernos da PCPE Agente
# (slug_disciplina, nome_exibicao, materia_tec, peso_prova, bloco_prova, caminho_edital, observacao_caderno)
CADERNOS_PCPE = [
    (
        "lingua-portuguesa",
        "Língua Portuguesa",
        "Língua Portuguesa (Português)",
        10,
        "P2 - Conhecimentos Específicos",
        "content/edital/pcpe/lingua-portuguesa.json",
        "Banca Cebraspe. Foco em compreensão/interpretação de textos, reescritura de frases, sintaxe de concordância/regência e Manual de Redação da Presidência da República (item 7 - Aula 13).",
    ),
    (
        "informatica",
        "Informática",
        "Informática",
        10,
        "P2 - Conhecimentos Específicos",
        "content/edital/pcpe/informatica.json",
        "Banca Cebraspe. Hardware e prompt (a14/c01) excluídos do edital da PCPE. Foco em segurança da informação, redes de computadores, computação em nuvem e ferramentas de escritório.",
    ),
    (
        "contabilidade-geral",
        "Contabilidade Geral",
        "Contabilidade Geral",
        10,
        "P2 - Conhecimentos Específicos",
        "content/edital/contabilidade-geral.json",
        "Banca Cebraspe. Disciplina nova. Foco em noções básicas, patrimônio, escrituração, lançamentos, Balanço Patrimonial e DRE. Aulas de CPC avançado e reservas são complementares.",
    ),
    (
        "direito-constitucional",
        "Direito Constitucional",
        "Direito Constitucional (CF/1988 e Doutrina)",
        4,
        "P1 - Noções de Direito",
        "content/edital/pcpe/direito-constitucional.json",
        "Banca Cebraspe. Excluídos: Ordem Social (a16), Controle de Constitucionalidade (a17) e complementos c01-c05. Foco em art. 5º (direitos e garantias), art. 37-41 (adm. pública) e art. 144 (segurança pública).",
    ),
    (
        "direito-administrativo",
        "Direito Administrativo",
        "Direito Administrativo (Doutrina e Leis Federais)",
        4,
        "P1 - Noções de Direito",
        "content/edital/pcpe/direito-administrativo.json",
        "Banca Cebraspe. Excluída Lei 8.112/90 pura (o regime do Agente é estadual - Lei 6.123/PE). Foco em atos administrativos, poderes, responsabilidade civil e Nova Lei de Licitações (Lei 14.133/2021 - a20 a a22).",
    ),
    (
        "legislacao-pcpe",
        "Legislação PCPE",
        "Legislação Estadual › Pernambuco (Legislação Estadual de PE)",
        4,
        "P1 - Noções de Direito",
        "content/edital/legislacao-pcpe.json",
        "Banca Cebraspe / outras bancas de PE. Disciplina nova. Estatuto da Polícia Civil de PE (Lei 6.425/1972), Lei Orgânica da PCPE (LC 424/2019) e Estatuto dos Servidores (Lei 6.123/1968). Se faltar Cebraspe, aceitar UPENET/IAUPE, FGV e FCC.",
    ),
    (
        "estatistica",
        "Estatística",
        "Estatística",
        4,
        "P2 - Conhecimentos Específicos",
        "content/edital/pcpe/estatistica.json",
        "Banca Cebraspe. Aulas a07 a a10 excluídas (distribuições contínuas/discretas fora do edital). Foco em estatística descritiva (médias, mediana, moda, dispersão, separatrizes, gráficos e correlação linear).",
    ),
    (
        "direito-penal",
        "Direito Penal",
        "Direito Penal",
        3,
        "P1 - Noções de Direito",
        "content/edital/direito-penal.json",
        "Banca Cebraspe. Disciplina nova. Foco em teoria do crime (fato típico, ilicitude, culpabilidade), crimes contra a pessoa, crimes contra o patrimônio e crimes contra a administração pública.",
    ),
    (
        "direito-processual-penal",
        "Direito Processual Penal",
        "Direito Processual Penal",
        3,
        "P1 - Noções de Direito",
        "content/edital/direito-processual-penal.json",
        "Banca Cebraspe. Disciplina nova. Foco prioritário em Inquérito Policial (a01), Provas (a03) e Prisão em Flagrante / Preventiva / Temporária (a04).",
    ),
    (
        "raciocinio-logico",
        "Raciocínio Lógico",
        "Raciocínio Lógico",
        3,
        "P2 - Conhecimentos Específicos",
        "content/edital/pcpe/raciocinio-logico.json",
        "Banca Cebraspe. Foco em lógica de proposições, conectivos, equivalências lógicas, leis de De Morgan, negações e lógica de argumentação.",
    ),
    (
        "matematica",
        "Matemática",
        "Matemática",
        3,
        "P2 - Conhecimentos Específicos",
        "content/edital/pcpe/matematica.json",
        "Banca Cebraspe. Foco em conjuntos numéricos, equações e sistemas lineares, contagem, probabilidade, progressões e matrizes. Geometria analítica e matemática financeira excluídas.",
    ),
    (
        "legislacao-penal-especial",
        "Legislação Penal Especial",
        "Legislação Penal e Processual Penal Especial",
        2,
        "P1 - Noções de Direito",
        "content/edital/legislacao-penal-especial.json",
        "Banca Cebraspe. Disciplina nova. Foco em Lei de Drogas (11.343/2006), Estatuto do Desarmamento (10.826/2003), Abuso de Autoridade (13.869/2019), Crimes Hediondos (8.072/1990) e Organizações Criminosas (12.850/2013).",
    ),
]


def consultar_aprovadas_cebraspe() -> dict[str, dict[str, int]]:
    """Consulta o banco de dados local via script em node para obter contagem
    de questões Cebraspe com status 'APPROVED', agrupadas por subjectId e editalItemId/aulaId.
    Retorna dicionário vazio caso o banco esteja indisponível ou ocorra erro.
    """
    counts: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    script_path = WEB / "scripts" / "questoes-contagem-cebraspe.ts"
    if not script_path.exists():
        return counts

    try:
        res = subprocess.run(
            ["npx", "tsx", "scripts/questoes-contagem-cebraspe.ts"],
            cwd=str(WEB),
            capture_output=True,
            text=True,
            encoding="utf-8",
            shell=sys.platform == "win32",
            timeout=15,
        )
        if res.returncode == 0 and "JSON_START" in res.stdout:
            start = res.stdout.index("JSON_START") + len("JSON_START")
            end = res.stdout.index("JSON_END", start)
            data = json.loads(res.stdout[start:end])
            for item in data:
                sub = item.get("subjectId", "")
                aula = item.get("aulaId", "")
                it = item.get("editalItemId", "")
                cnt = item.get("_count", {}).get("id", 0)
                if sub:
                    if it:
                        counts[sub][it] += cnt
                    if aula:
                        counts[sub][aula] += cnt
    except Exception:
        # Se falhar a conexão ao banco, segue com contagem zerada
        pass
    return counts


def carregar_topicos_edital(caminho_relativo: str) -> list[dict]:
    """Carrega os itens do edital de um arquivo JSON do projeto."""
    arquivo = ROOT / caminho_relativo
    if not arquivo.exists():
        return []
    try:
        dados = json.loads(arquivo.read_text(encoding="utf-8"))
        return dados.get("items", [])
    except Exception:
        return []


def criar_planilha_pcpe() -> Path:
    wb = Workbook()

    # Estilos
    cor_azul_escuro = "1F3A5F"
    cor_azul_claro = "EBF2FA"
    fonte_cabecalho = Font(bold=True, color="FFFFFF", size=11)
    fundo_cabecalho = PatternFill("solid", fgColor=cor_azul_escuro)
    alinhamento_cabecalho = Alignment(wrap_text=True, vertical="center", horizontal="center")
    alinhamento_padrao = Alignment(wrap_text=True, vertical="top")
    alinhamento_centro = Alignment(wrap_text=True, vertical="top", horizontal="center")
    alinhamento_direita = Alignment(wrap_text=True, vertical="top", horizontal="right")
    fonte_destaque = Font(bold=True)

    aprovadas_db = consultar_aprovadas_cebraspe()

    # ──────────────────────────────────────────────────────────────────────────
    # 1. ABA 'Cadernos'
    # ──────────────────────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Cadernos"

    cabecalhos_ws1 = [
        "Ordem",
        "Disciplina (Caderno)",
        "Bloco da Prova",
        "Questões na Prova",
        "Matéria no TecConcursos",
        "Banca",
        "Anos Recomendados",
        "Tópicos do Edital",
        "Cebraspe Aprovadas",
        "Meta de Questões",
        "Faltam Baixar",
        "Observações e Orientações Técnicas",
    ]
    larguras_ws1 = [8, 26, 26, 12, 36, 14, 18, 12, 12, 14, 14, 60]

    ws1.append(cabecalhos_ws1)
    for col_idx, largura in enumerate(larguras_ws1, start=1):
        ws1.column_dimensions[get_column_letter(col_idx)].width = largura
        cel = ws1.cell(1, col_idx)
        cel.font = fonte_cabecalho
        cel.fill = fundo_cabecalho
        cel.alignment = alinhamento_cabecalho
    ws1.freeze_panes = "A2"

    linhas_topicos = []
    tot_meta_global = 0
    tot_aprovadas_global = 0
    tot_deficit_global = 0

    for ordem, (slug, nome, mat_tec, peso, bloco, arq_edital, obs) in enumerate(CADERNOS_PCPE, start=1):
        topicos = carregar_topicos_edital(arq_edital)
        num_topicos = len(topicos)

        # Meta de mínimo 10 por tópico
        meta_caderno = num_topicos * 10
        aprovadas_caderno = sum(aprovadas_db.get(slug, {}).values())
        deficit_caderno = max(0, meta_caderno - aprovadas_caderno)

        tot_meta_global += meta_caderno
        tot_aprovadas_global += aprovadas_caderno
        tot_deficit_global += deficit_caderno

        ws1.append([
            ordem,
            nome,
            bloco,
            peso,
            mat_tec,
            "CEBRASPE",
            "2025–26 (prio) / 2022–24 (comp)",
            num_topicos,
            aprovadas_caderno,
            meta_caderno,
            deficit_caderno,
            obs,
        ])

        # Preparar dados para a aba detalhada de tópicos
        for it in topicos:
            item_id = it.get("id", "")
            item_texto = it.get("item", "")
            status_item = it.get("status", "covered")
            where_list = it.get("where", [])
            aulas_vinculadas = ", ".join(w.get("aula", "") for w in where_list if w.get("aula"))
            topicos_vinculados = " / ".join(w.get("topic", "") for w in where_list if w.get("topic"))

            aprovadas_item = aprovadas_db.get(slug, {}).get(f"{slug}/e{item_id}", 0)
            if aprovadas_item == 0 and aulas_vinculadas:
                for w in where_list:
                    a_id = w.get("aula")
                    if a_id and a_id in aprovadas_db.get(slug, {}):
                        aprovadas_item += aprovadas_db[slug][a_id]

            meta_item = 10
            deficit_item = max(0, meta_item - aprovadas_item)

            linhas_topicos.append({
                "ordem": ordem,
                "caderno": nome,
                "slug": slug,
                "item_id": item_id,
                "item_texto": item_texto,
                "status": status_item,
                "aulas": aulas_vinculadas or "Complementar/Lacuna",
                "topicos": topicos_vinculados or item_texto,
                "materia_tec": mat_tec,
                "aprovadas": aprovadas_item,
                "meta": meta_item,
                "deficit": deficit_item,
            })

    # Linha de total na aba 1
    linha_total = [
        "TOTAL",
        "12 Disciplinas PCPE",
        "60 Questões (P1: 20 | P2: 40)",
        60,
        "—",
        "CEBRASPE",
        "—",
        len(linhas_topicos),
        tot_aprovadas_global,
        tot_meta_global,
        tot_deficit_global,
        "Meta global calculada: mínimo de 10 questões Cebraspe por aula/tópico do edital.",
    ]
    ws1.append(linha_total)
    row_idx_total = ws1.max_row
    for col_idx in range(1, len(cabecalhos_ws1) + 1):
        cel = ws1.cell(row_idx_total, col_idx)
        cel.font = fonte_destaque
        cel.fill = PatternFill("solid", fgColor=cor_azul_claro)

    # Formatar alinhamentos da aba 1
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row):
        row[0].alignment = alinhamento_centro
        row[1].alignment = alinhamento_padrao
        row[2].alignment = alinhamento_padrao
        row[3].alignment = alinhamento_centro
        row[4].alignment = alinhamento_padrao
        row[5].alignment = alinhamento_centro
        row[6].alignment = alinhamento_centro
        row[7].alignment = alinhamento_centro
        row[8].alignment = alinhamento_centro
        row[9].alignment = alinhamento_centro
        row[10].alignment = alinhamento_centro
        row[11].alignment = alinhamento_padrao

    # ──────────────────────────────────────────────────────────────────────────
    # 2. ABA 'Tópicos do Edital'
    # ──────────────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Tópicos do Edital")

    cabecalhos_ws2 = [
        "Caderno",
        "Item Edital",
        "Tópico Oficial do Edital PCPE",
        "Aulas do Sistema",
        "Assunto TEC / Termo de Busca Recomendado",
        "Cebraspe Aprovadas",
        "Meta Mínima",
        "Faltam (Déficit)",
        "Status Cobertura",
    ]
    larguras_ws2 = [24, 12, 50, 24, 55, 12, 12, 12, 14]

    ws2.append(cabecalhos_ws2)
    for col_idx, largura in enumerate(larguras_ws2, start=1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = largura
        cel = ws2.cell(1, col_idx)
        cel.font = fonte_cabecalho
        cel.fill = fundo_cabecalho
        cel.alignment = alinhamento_cabecalho
    ws2.freeze_panes = "A2"

    for r in linhas_topicos:
        ws2.append([
            r["caderno"],
            f"Item {r['item_id']}",
            r["item_texto"],
            r["aulas"],
            f"{r['topicos'][:90]} [{r['materia_tec'].split(' (')[0]}]",
            r["aprovadas"],
            r["meta"],
            r["deficit"],
            r["status"].upper(),
        ])

    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        row[0].alignment = alinhamento_padrao
        row[1].alignment = alinhamento_centro
        row[2].alignment = alinhamento_padrao
        row[3].alignment = alinhamento_padrao
        row[4].alignment = alinhamento_padrao
        row[5].alignment = alinhamento_centro
        row[6].alignment = alinhamento_centro
        row[7].alignment = alinhamento_centro
        row[8].alignment = alinhamento_centro

    ws2.auto_filter.ref = ws2.dimensions

    # ──────────────────────────────────────────────────────────────────────────
    # 3. ABA 'Filtros e Instruções'
    # ──────────────────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Filtros e Instruções")
    cabecalhos_ws3 = ["Parâmetro no TecConcursos", "Valor / Instrução Recomendada"]
    larguras_ws3 = [32, 100]

    ws3.append(cabecalhos_ws3)
    for col_idx, largura in enumerate(larguras_ws3, start=1):
        ws3.column_dimensions[get_column_letter(col_idx)].width = largura
        cel = ws3.cell(1, col_idx)
        cel.font = fonte_cabecalho
        cel.fill = fundo_cabecalho
        cel.alignment = alinhamento_cabecalho
    ws3.freeze_panes = "A2"

    instrucoes = [
        ("Pasta no TecConcursos", "Questões Missão / PCPE Agente (criar na conta do TecConcursos)."),
        ("Banca Examinadora", "CEBRASPE (CESPE) — banca oficial do concurso da PCPE."),
        (
            "Anos de Prova (Filtro 1 - Prioritário)",
            "2025 e 2026. Priorizar questões recentes de múltipla escolha e estilo Cebraspe.",
        ),
        (
            "Anos de Prova (Filtro 2 - Complementar)",
            "2022 a 2024. Usar para disciplinas novas ou assuntos com poucas questões em 2025–2026 (ex.: Legislação PCPE, Contabilidade, Legislação Especial).",
        ),
        ("Escolaridade", "Nível Superior e Nível Médio (selecionar ambos)."),
        ("Opções de Filtragem", "Marcar estritamente: [X] Remover anuladas e [X] Remover desatualizadas."),
        (
            "Formato das Questões",
            "Múltipla Escolha (A, B, C, D, E). O edital da PCPE prevê prova objetiva de múltipla escolha.",
        ),
        (
            "Meta Quantitativa",
            "Mínimo de 10 questões por tópico/aula do edital (descontando as já aprovadas no banco).",
        ),
        (
            "Configuração de Impressão (PDF)",
            "Quantidade máxima de 200 questões por arquivo; Início 'a partir da questão' = última + 1; Cabeçalho com matéria e assunto; Gabarito ao final; Texto associado incluído; Sem rascunho; QR code não.",
        ),
        (
            "Atenção com Legislação PCPE",
            "Buscar no TEC por 'Legislação Estadual de Pernambuco' ou 'Polícia Civil de PE' (Leis 6.425/1972, 6.123/1968 e LC 424/2019). Caso haja poucas questões Cebraspe, aceitar bancas regionais (UPENET/IAUPE, FGV, FCC).",
        ),
        (
            "Atenção com Direito Administrativo",
            "Não incluir Lei 8.112/1990 pura (o regime do cargo é estadual). Priorizar Nova Lei de Licitações (Lei 14.133/2021), atos administrativos, poderes e responsabilidade civil.",
        ),
        (
            "Destino dos PDFs e Importação",
            "Salvar os PDFs gerados na pasta 'docs/questoes/' e executar o comando de importação local:\n"
            "pipeline/.venv/Scripts/python pipeline/questoes/importar.py --so-extrair\n"
            "E após conferir o relatório: pipeline/.venv/Scripts/python pipeline/questoes/importar.py",
        ),
    ]

    for param, val in instrucoes:
        ws3.append([param, val])

    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
        row[0].alignment = alinhamento_padrao
        row[1].alignment = alinhamento_padrao

    saida = ROOT / "docs/questoes/downloads/cadernos-a-montar-pcpe.xlsx"
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return saida


def main():
    arquivo_gerado = criar_planilha_pcpe()
    print(f"Sucesso: Planilha gerada em {arquivo_gerado}")


if __name__ == "__main__":
    main()
