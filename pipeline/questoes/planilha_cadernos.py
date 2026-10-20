"""Planilha dos cadernos a montar no TecConcursos (montagem manual pelo dono).

    pipeline/.venv/Scripts/python pipeline/questoes/planilha_cadernos.py

Entradas: docs/questoes/relatorios/prioridades-download.md (déficit por item do edital, gerado por prioridades.py),
docs/questoes/auditoria/2026-10-14/mapa_correspondencia.csv (assuntos do TEC já vistos nos PDFs → item do sistema)
e docs/questoes/downloads/cadernos.json. Saída: docs/questoes/downloads/cadernos-a-montar.xlsx
"""
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.stdout.reconfigure(encoding="utf-8")

# disciplina do sistema → (nome do caderno, matéria no TecConcursos, concursos, prioridade)
CAD = [
    ("legislacoes-pe", "Legislações PE", "Legislação Militar › Legislação Militar de PE (Estatuto 6.783/1974, Código Disciplinar 11.817/2000); Lei 15.187/2013, Decreto 50.014/2020 e Lei 14.751/2023 se aparecerem como assunto próprio", "CBMPE O/S, PMPE O/S", 1),
    ("direito-penal-militar", "Direito Penal Militar", "Direito Penal Militar", "CBMPE O, PMPE O", 1),
    ("direito-constitucional", "Direito Constitucional", "Direito Constitucional (CF/1988 e Doutrina)", "CBMPE O/S, PMPE O/S, PCPE", 1),
    ("fisica", "Física", "Física", "CBMPE O/S", 2),
    ("quimica", "Química", "Química e Engenharia Química", "CBMPE O", 2),
    ("matematica", "Matemática", "Matemática", "CBMPE O/S", 2),
    ("lingua-espanhola", "Língua Espanhola", "Língua Espanhola (Espanhol)", "CBMPE O, PMPE O", 2),
    ("biologia", "Biologia", "Biologia e Biomedicina", "CBMPE O/S", 3),
    ("lingua-inglesa", "Língua Inglesa", "Língua Inglesa (Inglês)", "CBMPE O, PMPE O", 3),
    ("direito-administrativo", "Direito Administrativo", "Direito Administrativo (Doutrina e Leis Federais)", "CBMPE O, PMPE O, PCPE", 3),
    ("informatica", "Informática", "Informática", "todos", 3),
    ("estatistica", "Estatística", "Estatística", "CBMPE O, PMPE O, PCPE", 3),
    ("lingua-portuguesa", "Língua Portuguesa", "Língua Portuguesa (Português)", "todos", 4),
    ("raciocinio-logico", "Raciocínio Lógico", "Raciocínio Lógico", "CBMPE S, PMPE O/S, PCPE", 4),
    ("atualidades", "Atualidades", "Atualidades e Conhecimentos Gerais (só assuntos atemporais)", "CBMPE S", 4),
    ("historia-pe", "História de Pernambuco", "História (procure Pernambuco / História de PE)", "CBMPE S, PMPE S", 4),
]
# disciplinas de PMPE/PCPE que ainda não têm aula no sistema: tópicos do edital como guia
EXTRA = [
    ("Direito Penal", "Direito Penal", "PMPE O, PCPE", 5, "Lei penal; teoria do crime; penas; extinção da punibilidade; crimes contra a pessoa, patrimônio, dignidade sexual, fé pública e administração pública; Leis 13.869 (abuso de autoridade), 7.716, 8.072, 9.503, 9.455, 9.605, 10.826, 11.340, 11.343, 14.344"),
    ("Direito Processual Penal", "Direito Processual Penal", "PMPE O, PCPE", 5, "Princípios; inquérito policial; ação penal; jurisdição e competência; provas; prisões e cautelares; liberdade provisória; procedimentos; nulidades; recursos; Leis 7.960, 9.099, 9.296, 12.850"),
    ("Direito Processual Penal Militar", "Direito Processual Penal Militar", "PMPE O", 5, "Aplicação da lei; polícia judiciária militar; inquérito policial militar; ação penal militar; competência; denúncia; provas; processo ordinário e especiais; nulidades e recursos"),
    ("Direitos e Garantias Fundamentais / Direitos Humanos", "Direitos Humanos (e Estatutos: ECA, Idoso, Pessoa com Deficiência, Juventude)", "PMPE O/S", 5, "DUDH; Pacto de San José; CEDAW; Convenção sobre os Direitos da Criança; ECA (8.069/1990); Estatuto da Pessoa Idosa (10.741/2003); Estatuto da Pessoa com Deficiência (13.146/2015); Estatuto da Juventude (12.852/2013); tortura; uso da força (13.060/2014)"),
    ("Legislação Penal Especial", "Legislação Penal e Processual Penal Especial", "PCPE", 5, "Leis especiais penais conforme o edital da PCPE"),
    ("Contabilidade Geral", "Contabilidade Geral", "PCPE Agente", 5, "Conforme edital da PCPE"),
    ("Arquivologia", "Arquivologia", "PCPE Escrivão", 5, "Conforme edital da PCPE"),
]

# --- déficit por item, a partir do relatório de prioridades
txt = (ROOT / "docs/questoes/relatorios/prioridades-download.md").read_text(encoding="utf-8")
defic = defaultdict(dict)  # disc -> item_code -> (titulo, aprovadas)
cur = None
for ln in txt.splitlines():
    m = re.match(r"## ([a-z-]+): ", ln)
    if m:
        cur = m.group(1)
        continue
    m = re.match(r"- `([^`]+)` (.+?)(?:: (\d+) de 10)?$", ln)
    if m and cur:
        defic[cur][m.group(1)] = (m.group(2), int(m.group(3) or 0))

# --- assuntos do TEC já vistos nos PDFs, ligados a itens do sistema
mp = list(csv.DictReader(open(ROOT / "docs/questoes/auditoria/2026-10-14/mapa_correspondencia.csv", encoding="utf-8-sig"), delimiter=";"))
tec_por_item = defaultdict(list)  # (disc,item) -> [(materia, assunto, questoes)]
for r in mp:
    if r["destinoDisciplina"] and r["destinoItensEdital"] and r["tipo"] in ("DIRETA", "ALIAS", "PARCIAL"):
        for it in [x.strip() for x in r["destinoItensEdital"].split(",") if x.strip()]:
            d, _, code = it.partition("/")
            tec_por_item[(d, code)].append((r["materiaTEC"], r["assuntoTEC"], int(r["questoes"])))

wb = Workbook()
H = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3A5F")
wrap = Alignment(wrap_text=True, vertical="top")


def sheet(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i)
        c.font, c.fill, c.alignment = H, HF, Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


# 1) Cadernos
ws = wb.active
ws.title = "Cadernos"
sheet(ws, ["Ordem", "Nome do caderno", "Matéria no TecConcursos", "Concursos", "Prioridade", "Assuntos a marcar", "Questões que faltam (meta 10 por assunto)", "Observação"], [7, 28, 48, 20, 10, 12, 18, 60])
ordem = 0
obs = {
    "legislacoes-pe": "Só 4 questões em 2025–26 com os filtros comuns. Se vier pouco, aceite anos anteriores (decisão sua) ou busque outras bancas.",
    "fisica": "Já existe o 'Caderno Física' (236 questões) na pasta Questões Missão. Conferir antes de refazer.",
    "direito-administrativo": "Evitar licitações, improbidade e LAI (já há ~1.100 questões guardadas sem uso).",
    "atualidades": "Só assuntos atemporais. Evitar 'Realidade ... de outros estados' e fatos datados.",
    "lingua-portuguesa": "Fonética e Separação Silábica fora (decisão do dono).",
    "historia-pe": "No TEC pode ser matéria 'História'; filtrar assuntos de Pernambuco.",
}
n_ass = {}
for disc, nome, mat, conc, prio in sorted(CAD, key=lambda x: x[4]):
    items = defic.get(disc, {})
    falta = sum(10 - n for _, n in items.values())
    ordem += 1
    n_ass[disc] = len(items)
    ws.append([ordem, nome, mat, conc, prio, len(items), falta, obs.get(disc, "")])
for nome, mat, conc, prio, tops in EXTRA:
    ordem += 1
    ws.append([ordem, nome, mat, conc, prio, "—", "—", "Sem aula no sistema ainda (edital PMPE/PCPE). Marque os assuntos conforme o edital: " + tops])
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = wrap

# 2) Assuntos por caderno
ws2 = wb.create_sheet("Assuntos por caderno")
sheet(ws2, ["Caderno", "Item do edital (sistema)", "Assunto do edital", "Aprovadas hoje", "Faltam", "Assuntos do TecConcursos a marcar (nome exato visto nos PDFs)", "Se não houver assunto equivalente, busque por"], [24, 12, 60, 11, 8, 80, 40])
for disc, nome, mat, conc, prio in sorted(CAD, key=lambda x: x[4]):
    items = defic.get(disc, {})
    for code, (tit, n) in sorted(items.items(), key=lambda kv: (kv[1][1], kv[0])):
        tec = tec_por_item.get((disc, code), [])
        tecs = "; ".join(f"{a} [{m.split(' (')[0]}]" for m, a, q in sorted(tec, key=lambda x: -x[2])[:6]) if tec else ""
        ws2.append([nome, code, tit, n, 10 - n, tecs, "" if tecs else tit[:80]])
# disciplinas fora do CBMPE Oficial (sem déficit calculado): lista os assuntos do TEC já vistos nos PDFs
for disc, nome, mat, conc, prio in sorted(CAD, key=lambda x: x[4]):
    if defic.get(disc):
        continue
    vistos = [r for r in mp if r["destinoDisciplina"] == disc and r["tipo"] in ("DIRETA", "ALIAS", "PARCIAL")]
    for r in sorted(vistos, key=lambda r: -int(r["questoes"])):
        ws2.append([nome, r["destinoItensEdital"].replace(disc + "/", ""), "(assunto do TEC visto nos PDFs; " + r["questoes"] + " questões já baixadas)", "", "", f'{r["assuntoTEC"]} [{r["materiaTEC"].split(" (")[0]}]', ""])
for row in ws2.iter_rows(min_row=2):
    for c in row:
        c.alignment = wrap
ws2.auto_filter.ref = ws2.dimensions

# 3) Filtros e impressão
ws3 = wb.create_sheet("Filtros e impressão")
sheet(ws3, ["Item", "Valor"], [34, 110])
for k, v in [
    ("Pasta", "Questões Missão (já criada na sua conta)"),
    ("Bancas", "AOCP, Instituto AOCP (prioridade); CEBRASPE (CESPE), FGV, FCC (complemento quando faltar AOCP)"),
    ("Anos", "2025 e 2026"),
    ("Opções", "Remover anuladas; Remover desatualizadas"),
    ("Dica", "Marque Banca, Ano e Opções primeiro e vá trocando só os assuntos de cada caderno. 'Salvar filtros' do site não funcionou (retornou 'Seleção não foi salva')."),
    ("Atenção", "Use os assuntos específicos, não 'Todo o conteúdo' da matéria. A matéria inteira traz fora do edital (técnicas de laboratório, outros estados etc.)."),
    ("Impressão", "Quantidade máxima 200; início 'a partir da questão' = última impressa + 1; cabeçalho com matéria e assunto; gabarito no fim; fonte normal; texto associado sim; QR code não; sem rascunho; remover nenhuma."),
    ("Limite diário", "10 PDFs por dia (200 questões cada). Registro de progresso: docs/questoes/downloads/cadernos.json"),
    ("Depois de baixar", "Colocar os PDFs em docs/questoes/ e rodar: pipeline/.venv/Scripts/python pipeline/questoes/importar.py --so-extrair"),
    ("Aviso", "Com banca e ano restritos, assuntos de nicho podem trazer poucas questões (Legislação Militar de PE: 4). Se o caderno vier pequeno, avise para ampliarmos os anos."),
]:
    ws3.append([k, v])
for row in ws3.iter_rows(min_row=2):
    for c in row:
        c.alignment = wrap

out = ROOT / "docs/questoes/downloads/cadernos-a-montar.xlsx"
wb.save(out)
print(out, "| cadernos:", ordem, "| linhas de assunto:", ws2.max_row - 1)
